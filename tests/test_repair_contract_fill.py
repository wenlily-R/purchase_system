#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""回归测试: 维修订单生成《修理修缮合同》全字段自动填充 (V11.300)

需求(用户《系统修改》): 维修订单生成合同时自动读取维修订单数据, 填充合同全部空白区域 ——
  ① 合同头部: 乙方名称 / 委托修理标的物名称
  ② 《修理修缮物项目清单及价款》表格: 按订单明细逐行填 品名/规格型号/计量单位/数量/单价/金额,
     行数随明细动态增减; 合计金额/大写金额/税金/不含税价款(税率13%) 自动计算
  ③ 修理修缮时间、地点: 起止年月日自动填充(地点保持模板原文)
  ④ 第七条结算底部收款账户信息: 收款账户名称/收款账号/收款银行/银行行号 ← 供应商档案
  ⑤ 附《维修明细单》: 维修时间列 + 明细逐行 + 合计, 行数随明细动态增减
  ⑥ 不破坏既有: 买卖合同模板仍走原"三段式"合计句 + 序号表头明细表(含合计列口径)

跑法: .venv/Scripts/python.exe tests/test_repair_contract_fill.py   (Mac: python3 tests/test_repair_contract_fill.py)
全程 tempfile 副本库 + 打桩审批/消息发起(create_approvals/start_instances), 真实库零写入、零外部推送。
"""
import os
import sys
import json
import shutil
import sqlite3
import tempfile

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIVE = os.path.join(BASE, 'data', 'purchase.db')
SUP = '测试维修商V11300'
SUP2 = '测试物资商V11300'
P, F = [], []


def ck(n, c, e=''):
    (P if c else F).append(n)
    print(('  OK   ' if c else '  FAIL ') + n + (('  | ' + str(e)[:300]) if e else ''))


def fmt2(v):
    return '{:,.2f}'.format(round(float(v), 2))


def mk_order(c, tag, req_type, device, supplier, items, target_date='2026-09-20',
             created='2026-09-13 09:00:00'):
    """造夹具: 采购申请 + 订单 + 订单明细(单价/数量自算金额)"""
    c.execute("""INSERT INTO purchase_requests(req_no,dept,requester,requester_id,budget_code,purpose,target_date,
        status,total_estimated,remark,attachments,urgent,apply_date,req_type,repair_device,repair_fault)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              ('WX-' + tag, '维修车间', '穆娇', 1, '', '测试维修事由', target_date,
               '已通过', 0, '', '[]', 0, '2026-09-13', req_type, device, '皮带跑偏/异响'))
    rid = c.execute("SELECT id FROM purchase_requests WHERE req_no=?", ('WX-' + tag,)).fetchone()[0]
    tot = round(sum(float(i[3]) * float(i[4]) for i in items), 2)
    c.execute("""INSERT INTO purchase_orders(order_no,req_id,item_name,spec,quantity,unit,price,amount,tax_rate,
        tax_amount,total_amount,supplier,requester,category,owner,owner_id,target_date,status,created_at,updated_at,
        trade_mode,urgent,attachments,settle_type,freight)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              ('CG-' + tag, rid, items[0][0], items[0][1], items[0][3], items[0][2], items[0][4], tot, 13.0,
               0, tot, supplier, '穆娇', 'WXHT', '穆娇', 1, target_date, '已通过', created, created,
               '货到付款', 0, '[]', '现结', 0))
    oid = c.execute("SELECT id FROM purchase_orders WHERE order_no=?", ('CG-' + tag,)).fetchone()[0]
    for it in items:
        c.execute("""INSERT INTO order_items(order_id,item_name,spec,unit,quantity,price,amount,tax_rate,
            tax_amount,total_amount,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                  (oid, it[0], it[1], it[2], it[3], it[4], round(float(it[3]) * float(it[4]), 2), 13.0, 0,
                   round(float(it[3]) * float(it[4]), 2), '', created))
    c.commit()
    return oid


def rows_of(tb):
    return [[c.text.strip() for c in r.cells] for r in tb.rows]


def font_stat(doc):
    """非空 run 的 (ascii, eastAsia) 字体分布(段落+表格) —— 字体口径回归用"""
    from collections import Counter
    from docx.oxml.ns import qn
    cnt = Counter()

    def grab(paras):
        for pp in paras:
            for r in pp.runs:
                if not r.text.strip():
                    continue
                rpr = r._element.find(qn('w:rPr'))
                rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
                cnt[(rf.get(qn('w:ascii')) if rf is not None else None,
                     rf.get(qn('w:eastAsia')) if rf is not None else None)] += 1
    grab(doc.paragraphs)
    for tb in doc.tables:
        for rw in tb.rows:
            for cc in rw.cells:
                grab(cc.paragraphs)
    return cnt


def font_ok(stat):
    """中文一律仿宋 + 西文/数字仅 Times New Roman/仿宋(用户 2026-09-13 字体归一要求)"""
    return all(ea == '仿宋' for _, ea in stat) and all(a in ('仿宋', 'Times New Roman') for a, _ in stat)


def eff_ea(doc, run, para):
    """run 的**有效**中文字体: run rFonts → 段落样式链(Normal 等) → docDefaults
    (渲染时新建的 run 常无 rFonts, 靠模板 Normal 样式兜底, 只查属性会误报)"""
    from docx.oxml.ns import qn
    rpr = run._element.find(qn('w:rPr'))
    rf = rpr.find(qn('w:rFonts')) if rpr is not None else None
    if rf is not None and rf.get(qn('w:eastAsia')):
        return rf.get(qn('w:eastAsia'))
    st, seen = para.style, set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        srpr = st.element.find(qn('w:rPr'))
        srf = srpr.find(qn('w:rFonts')) if srpr is not None else None
        if srf is not None and srf.get(qn('w:eastAsia')):
            return srf.get(qn('w:eastAsia'))
        st = st.base_style
    dd = doc.styles.element.find(qn('w:docDefaults'))
    if dd is not None:
        rp = dd.find(qn('w:rPrDefault'))
        rp = rp.find(qn('w:rPr')) if rp is not None else None
        rf2 = rp.find(qn('w:rFonts')) if rp is not None else None
        if rf2 is not None:
            return rf2.get(qn('w:eastAsia'))
    return None


def non_fangsong_runs(doc):
    """全篇(段落+表格)非空 run 里, 有效中文字体不是仿宋的清单 —— 空清单=整篇中文渲染为仿宋"""
    bad = []

    def grab(paras):
        for pp in paras:
            for r in pp.runs:
                if r.text.strip() and eff_ea(doc, r, pp) != '仿宋':
                    bad.append((r.text[:24], eff_ea(doc, r, pp)))
    grab(doc.paragraphs)
    for tb in doc.tables:
        for rw in tb.rows:
            for cc in rw.cells:
                grab(cc.paragraphs)
    return bad


def main():
    import docx                      # noqa: 仅测试用
    tmp = tempfile.mkdtemp(prefix='test_repair_contract-')
    sys.path.insert(0, BASE)
    import app as A

    made_files = []
    live = sqlite3.connect(LIVE)
    live.row_factory = sqlite3.Row
    nm0 = live.execute("SELECT COUNT(*) c FROM contracts").fetchone()['c']
    live.close()

    copy = os.path.join(tmp, 'live.db')
    shutil.copy(LIVE, copy)
    A.DB = copy
    A.init_db()
    # 打桩: 不建审批实例、不发起钉钉/飞书推送(测试单绝不惊动领导手机)
    A.create_approvals = lambda *a, **k: None
    A.start_instances = lambda *a, **k: None

    c = sqlite3.connect(copy)
    c.row_factory = sqlite3.Row
    print('[1] 结构: 供应商档案 银行行号 列')
    cols = [r['name'] for r in c.execute("PRAGMA table_info(suppliers)").fetchall()]
    ck('suppliers.bank_no 已存在(迁移/init_db 双写)', 'bank_no' in cols, cols)
    _tf = font_stat(docx.Document(os.path.join(BASE, 'contract_templates', '修理修缮合同.docx')))
    ck('模板: 中文全仿宋、无宋体残留(数字/字母 Times New Roman)',
       font_ok(_tf) and not any(ea == '宋体' for _, ea in _tf), dict(_tf))
    c.execute("""INSERT INTO suppliers(name,contact,phone,category,level,bank,account,bank_no,tax_id,invoice_type,rating,status)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
              (SUP, '赵工', '13800000000', '维修类', '核心供应商', '中国工商银行河曲支行', '6222020200001111',
               '102168000123', '91140900MA0TEST01', '增值税专用发票', 4.8, '正常'))
    c.execute("""INSERT INTO suppliers(name,contact,phone,category,level,bank,account,bank_no,tax_id,invoice_type,rating,status)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
              (SUP2, '钱经理', '13900000000', '设备类', '一般供应商', '中国建设银行太原分行', '6217000011112222',
               '105161000012', '', '增值税专用发票', 4.0, '正常'))
    c.commit()

    print('[2] 多行明细: 《修理修缮合同》全字段自动填充')
    R_ITEMS = [['更换托辊', 'TD75-Φ89', '根', 4, 120.0],
               ['皮带修补', 'EP200 800mm', '处', 2, 350.0],
               ['电机轴承更换', 'Y2-132', '台', 1, 820.0]]
    oid_rep = mk_order(c, 'REP1', '设备维修', '1号皮带机', SUP, R_ITEMS)
    cli = A.app.test_client()
    ck('登录(mujiao)', (cli.post('/api/login', json={'username': 'mujiao', 'password': '123456'})
                        .get_json() or {}).get('success'))
    r = cli.post('/api/contracts/generate', json={'order_id': oid_rep, 'template_name': '修理修缮合同',
                                                  'category': 'XLXS', 'settle_type': '现结'})
    j = r.get_json() or {}
    ck('生成成功(HTTP 200)', r.status_code == 200 and j.get('success'), j)
    if not j.get('success'):
        _finish(tmp, made_files, nm0)
        return
    fpath = os.path.join(BASE, 'uploads', j['file'])
    made_files.append(j['file'])
    ck('合同文件已落盘', os.path.exists(fpath), fpath)
    d = docx.Document(fpath)
    txt = '\n'.join(p.text for p in d.paragraphs)
    _gf = font_stat(d)
    ck('生成件: 全篇中文渲染为仿宋(逐run有效字体: rFonts→样式→docDefaults, 无宋体兜底)',
       not non_fangsong_runs(d) and not any(ea == '宋体' for _, ea in _gf), dict(_gf))
    allt = txt + '\n' + '\n'.join(cc.text for tb in d.tables for rw in tb.rows for cc in rw.cells)  # 含表格单元格
    t0, t1 = rows_of(d.tables[0]), rows_of(d.tables[1])
    total = 2000.0
    amt, tax = round(total / 1.13, 2), round(total - round(total / 1.13, 2), 2)

    ck('① 头部 乙方名称自动带出', '乙方：' + SUP in txt, [x for x in txt.split('\n') if '乙方' in x][:2])
    ck('① 头部 委托修理标的物名称自动带出', '就甲方委托乙方修理1号皮带机相关事宜' in txt)
    ck('③ 修理修缮时间起止年月日自动填充',
       ('1、2026年09月13日至2026年09月20日') in txt, [x for x in txt.split('\n') if '日至' in x][:2])
    ck('③ 地点保持模板原文(不出现空白)', '地点：河曲县正成洗选煤有限责任公司' in txt)
    ck('② 项目清单表 数据行=明细数(1行模板→3行动态增行)', len(t0) == len(R_ITEMS) + 3, len(t0))
    ck('② 明细逐列填充(品名/规格/单位/数量/单价/金额)',
       t0[1] == ['更换托辊', 'TD75-Φ89', '根', '4', fmt2(120), fmt2(480)] and
       t0[2][0] == '皮带修补' and t0[3][0] == '电机轴承更换', t0[1:4])
    ck('② 合计行金额写"合计"右格(不再覆盖标签)', t0[-2][4] == '合计' and t0[-2][5] == fmt2(total), t0[-2])
    ck('② 合计金额/大写自动计算', ('合计金额：¥%s元' % fmt2(total)) in allt and
       ('（人民币大写金额：人民币%s）' % A.rmb_upper(total)) in allt,
       [x for x in allt.split('\n') if '合计金额' in x][:1])
    ck('② 税金/不含税/税率13%%自动计算',
       ('税金（税率 13 %%）为：¥%s元' % fmt2(tax)) in allt and ('不含税价款为：¥%s元' % fmt2(amt)) in allt and
       ('人民币%s' % A.rmb_upper(amt)) in allt,
       '税金=%s 不含税=%s' % (fmt2(tax), fmt2(amt)))
    ck('② 模板下划线空白全部填满(旧空白残留=0)',
       '¥            元' not in allt and '人民币大写金额：  ' not in allt)
    ck('④ 收款账户信息四项读供应商档案',
       ('收款账户名称：' + SUP) in allt and '收款账号：6222020200001111' in allt and
       '收款银行：中国工商银行河曲支行' in allt and '银行行号：102168000123' in allt,
       [x for x in allt.split('\n') if '收款' in x or '银行' in x])
    ck('⑤ 《维修明细单》表头原样', t1[0] == ['维修时间', '修理修缮项目品名', '规格型号', '计量单位', '数量', '单价', '金额'], t1[0])
    ck('⑤ 维修明细单 行数随明细增减(6空行→3行)',
       len(t1) == 1 + 3 + 1 and t1[1][1] == '更换托辊' and t1[3][5] == fmt2(820), t1)
    ck('⑤ 维修时间列自动填充', all(r[0] == '2026.09.13至2026.09.20' for r in t1[1:4]), [r[0] for r in t1[1:4]])
    ck('⑤ 维修明细单合计', t1[-1][5] == '合计' and t1[-1][6] == fmt2(total), t1[-1])
    ct = c.execute("SELECT * FROM contracts WHERE order_id=?", (oid_rep,)).fetchone()
    ck('合同落库: 名称=设备名维修合同 / 模板名 / 金额', ct['contract_name'] == '1号皮带机维修合同' and
       ct['template_name'] == '修理修缮合同' and round(ct['amount'], 2) == total, dict(ct) if ct else None)
    ck('合同编号类目=XLXS(修理修缮)', '-XLXS-' in (ct['contract_no'] if ct else ''), ct['contract_no'] if ct else '')

    print('[3] 单行明细: 动态减行')
    oid_one = mk_order(c, 'REP2', '设备维修', '3号空压机', SUP, [['更换空滤芯', 'LU75-8', '个', 2, 1600.0]])
    r2 = cli.post('/api/contracts/generate', json={'order_id': oid_one, 'template_name': '修理修缮合同'})
    j2 = r2.get_json() or {}
    ck('单行明细 生成成功', j2.get('success'), j2)
    if j2.get('success'):
        made_files.append(j2['file'])
        d2 = docx.Document(os.path.join(BASE, 'uploads', j2['file']))
        s0, s1 = rows_of(d2.tables[0]), rows_of(d2.tables[1])
        ck('① 项目清单表 预置空行清空→仅1行明细', len(s0) == 1 + 3 and s0[1][0] == '更换空滤芯', s0)
        ck('② 合计=单价×数量=3,200.00', s0[-2][5] == fmt2(3200), s0[-2])
        ck('⑤ 维修明细单 6预置行→1行', len(s1) == 1 + 1 + 1 and s1[-1][6] == fmt2(3200), s1)
        t2 = '\n'.join(p.text for p in d2.paragraphs)
        ck('③ 单行也填起止时间(订单无完工日→按7天兜底)',
           ('1、2026年09月13日至2026年09月20日') in t2, [x for x in t2.split('\n') if '日至' in x][:1])

    print('[4] 既有模板回归: 买卖合同仍为三段式 + 序号表头明细表')
    W_ITEMS = [['输送带', 'EP200', '米', 10, 200.0], ['托辊', 'TD75', '根', 5, 120.0]]
    oid_w = mk_order(c, 'WZ1', '物资采购', '', SUP2, W_ITEMS)
    r3 = cli.post('/api/contracts/generate', json={'order_id': oid_w, 'template_name': '买卖合同-现结'})
    j3 = r3.get_json() or {}
    ck('物资订单生成成功', j3.get('success'), j3)
    if j3.get('success'):
        made_files.append(j3['file'])
        d3 = docx.Document(os.path.join(BASE, 'uploads', j3['file']))
        w0 = rows_of(d3.tables[0])
        t3 = '\n'.join(p.text for p in d3.paragraphs)
        t3 += '\n' + '\n'.join(cc.text for tb in d3.tables for rw in tb.rows for cc in rw.cells)
        ck('买卖合同 明细表: 序号列+2行(删除多余空行)+合计写金额列(倒数第2列)',
           w0[0][0] == '序号' and len(w0) == 2 + 4 and w0[1][0] == '1' and w0[2][0] == '2' and
           w0[1 + 2][-2] == fmt2(2600), w0)
        ck('买卖合同 合计句仍为原三段式(未走修理修缮就地填充分支)',
           ('合计金额：¥%s元（大写金额：人民币' % fmt2(2600)) in t3 and '税金（税率 13%）为：' in t3,
           [x for x in t3.split('\n') if '合计金额' in x][:1])
        ck('买卖合同 收款信息仍自动填充(逐run分支兼容)', ('收款账号名称：' + SUP2) in t3)

    _finish(tmp, made_files, nm0, c)
    c.close()


def _finish(tmp, made_files, nm0, c=None):
    print('[5] 清理与复核')
    for f in made_files:
        p = os.path.join(BASE, 'uploads', f)
        if os.path.exists(p):
            os.remove(p)
    ck('测试生成的合同文件已清理', all(not os.path.exists(os.path.join(BASE, 'uploads', f)) for f in made_files))
    shutil.rmtree(tmp, ignore_errors=True)
    live = sqlite3.connect(LIVE)
    live.row_factory = sqlite3.Row
    nm1 = live.execute("SELECT COUNT(*) c FROM contracts").fetchone()['c']
    sup_live = live.execute("SELECT COUNT(*) c FROM suppliers WHERE name IN (?,?)", (SUP, SUP2)).fetchone()['c']
    live.close()
    ck('真实库 contracts 计数未变', nm0 == nm1, '%s → %s' % (nm0, nm1))
    ck('真实库未残留测试供应商', sup_live == 0, sup_live)
    print('\n===== 结果: %d 通过, %d 失败 =====' % (len(P), len(F)))
    for x in F:
        print('  FAIL:', x)
    sys.exit(1 if F else 0)


if __name__ == '__main__':
    main()
